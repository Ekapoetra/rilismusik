"""Offline regressions for legacy stream visibility, royalty reports and chat reads."""
import io
import unittest
from datetime import datetime, timedelta, timezone
from unittest.mock import patch

import test_vercel_preview  # Configures offline environment before real router imports.
from fastapi import HTTPException
from mongomock_motor import AsyncMongoMockClient
import storage_service
from routes import chat, label_analytics, royalty, royalty_report_export as report_export


class LegacyStreamVisibilityTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.db = AsyncMongoMockClient().test
        self.patches = [patch.object(label_analytics, "db_bg", self.db)]
        for item in self.patches:
            item.start()
        self.addCleanup(lambda: [item.stop() for item in self.patches])
        self.user = {"id": "u1", "role": "label", "status": "active"}
        await self.db.labels.insert_many([
            {"id": "l1", "user_id": "u1", "label_name": "Label Satu"},
            {"id": "other", "user_id": "u2", "label_name": "Other"},
        ])
        await self.db.royalty_imports.insert_many([
            {"id": "imp-legacy", "status": "dana_received"},
            {"id": "imp-new", "status": "published"},
            {"id": "imp-draft", "status": "draft"},
        ])
        base = {"match_status": "matched", "platform": "Spotify", "country": "ID",
                "track_title_raw": "Lagu", "artist_name_raw": "Artis"}
        await self.db.royalty_lines.insert_many([
            # Legacy cut-off: streams are the label's own, royalty stays hidden.
            {**base, "id": "legacy", "label_id": "l1", "import_id": "imp-legacy", "period": "2026-04",
             "status": "withdrawn", "legacy_settled": True, "quantity": 700, "label_idr": 70_000},
            {**base, "id": "new", "label_id": "l1", "import_id": "imp-new", "period": "2026-05",
             "status": "withdrawn", "legacy_settled": False, "quantity": 100, "label_idr": 1_000},
            # Legacy rows of an unpublished import must not appear yet.
            {**base, "id": "legacy-draft", "label_id": "l1", "import_id": "imp-draft", "period": "2026-06",
             "status": "withdrawn", "legacy_settled": True, "quantity": 9_999, "label_idr": 9_999},
            {**base, "id": "draft", "label_id": "l1", "import_id": "imp-draft", "period": "2026-06",
             "status": "draft", "quantity": 5_000, "label_idr": 5_000},
            {**base, "id": "staged", "label_id": "l1", "import_id": "imp-new", "period": "2026-06",
             "status": "pending", "replacement_stage": True, "quantity": 4_000, "label_idr": 4_000},
            {**base, "id": "foreign", "label_id": "other", "import_id": "imp-new", "period": "2026-05",
             "status": "available", "quantity": 3_000, "label_idr": 3_000},
        ])

    async def analytics(self, window):
        with patch.object(label_analytics, "get_labels_for_user", return_value=[{"id": "l1", "label_name": "Label Satu"}]), \
                patch.object(label_analytics, "account_entitlements", return_value={"multi_label": False}), \
                patch.object(label_analytics, "get_label_by_user", return_value={"id": "l1"}):
            return await label_analytics.get_label_dashboard_analytics(window=window, label_id=None, user=self.user)

    async def test_legacy_streams_are_visible_but_royalty_is_hidden(self):
        result = await self.analytics("6")
        self.assertEqual(result["latest_period"], "2026-05")
        self.assertEqual(result["totals"], {"streams": 800, "revenue_idr": 1_000, "lines": 2})
        by_period = {row["period"]: row for row in result["monthly"]}
        self.assertEqual(by_period["2026-04"]["streams"], 700)
        self.assertEqual(by_period["2026-04"]["revenue_idr"], 0)
        self.assertTrue(by_period["2026-04"]["revenue_hidden"])
        self.assertFalse(by_period["2026-05"]["revenue_hidden"])
        self.assertEqual(result["legacy"]["streams"], 700)
        self.assertEqual(result["legacy"]["periods"], ["2026-04"])
        self.assertEqual(result["top_tracks"][0]["revenue_idr"], 1_000)
        self.assertEqual(sorted(result["available_periods"]), ["2026-04", "2026-05"])

    async def test_label_with_only_legacy_history_still_gets_stream_chart(self):
        await self.db.royalty_lines.delete_many({"id": "new"})
        result = await self.analytics("latest")
        self.assertEqual(result["latest_period"], "2026-04")
        self.assertEqual(result["latest_report"], {"streams": 700, "revenue_idr": 0, "lines": 1})
        self.assertTrue(result["latest_report_revenue_hidden"])
        self.assertEqual([row["streams"] for row in result["monthly"]], [700])


class RoyaltyReportExportTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.db = AsyncMongoMockClient().test
        self.objects = {}

        async def upload(*, key, data, content_type):
            self.objects[key] = (data, content_type)
            return key

        async def presign(*, key, ttl, filename=None):
            return f"https://r2.example.invalid/{key}?ttl={ttl}&name={filename}"

        self.patches = [
            patch.object(report_export, "db_bg", self.db), patch.object(royalty, "db_bg", self.db),
            patch.object(royalty, "get_label_by_user", return_value={"id": "l1"}),
            patch.object(storage_service, "is_configured", return_value=True),
            patch.object(storage_service, "upload_bytes", side_effect=upload),
            patch.object(storage_service, "generate_presigned_url", side_effect=presign),
        ]
        for item in self.patches:
            item.start()
        self.addCleanup(lambda: [item.stop() for item in self.patches])
        self.user = {"id": "u1", "role": "label", "status": "active"}
        base = {"label_id": "l1", "platform": "Spotify", "country": "ID", "isrc": "ID0000000001"}
        await self.db.royalty_lines.insert_many([
            {**base, "id": "legacy", "period": "2025-12", "status": "withdrawn", "legacy_settled": True,
             "artist_name_raw": "Artis A", "track_title_raw": "Legacy Hidden", "quantity": 9, "label_idr": 9_999},
            {**base, "id": "paid", "period": "2026-02", "status": "withdrawn", "legacy_settled": False,
             "artist_name_raw": "Artis A", "track_title_raw": "Sudah Cair", "quantity": 10, "label_idr": 1_000},
            {**base, "id": "open", "period": "2026-03", "status": "available",
             "artist_name_raw": "Artis B", "track_title_raw": "Tersedia", "quantity": 20, "label_idr": 2_000},
            {**base, "id": "draft", "period": "2026-03", "status": "draft",
             "artist_name_raw": "Artis B", "track_title_raw": "Draft Hidden", "quantity": 99, "label_idr": 99},
            {**base, "id": "foreign", "label_id": "other", "period": "2026-03", "status": "available",
             "artist_name_raw": "Artis C", "track_title_raw": "Foreign Hidden", "quantity": 1, "label_idr": 1},
        ])

    async def test_csv_includes_withdrawn_lines_after_legacy_cutoff(self):
        base = await royalty._royalty_report_scope(self.user)
        text = (await report_export.build_csv_bytes(base)).decode()
        self.assertTrue(text.startswith("period,release_title,track_title"))
        self.assertIn("Sudah Cair", text)
        self.assertIn("Sudah Dicairkan,Ya", text)
        self.assertIn("Tersedia", text)
        for hidden in ("Legacy Hidden", "Draft Hidden", "Foreign Hidden"):
            self.assertNotIn(hidden, text)

    async def test_workbook_has_summary_and_artist_sheets(self):
        from openpyxl import load_workbook
        base = await royalty._royalty_report_scope(self.user)
        data = await royalty._build_royalty_workbook_bytes(base=base, period=None, artist=None)
        wb = load_workbook(io.BytesIO(data))
        self.assertEqual(wb.sheetnames, ["Ringkasan", "Artis A", "Artis B"])
        self.assertEqual(wb["Artis A"]["C2"].value, "Sudah Cair")
        self.assertEqual(wb["Artis A"]["L2"].value, "Ya")
        self.assertEqual(wb["Artis A"].max_row, 2)
        summary_values = [cell for row in wb["Ringkasan"].iter_rows(values_only=True) for cell in row]
        self.assertIn(3_000, summary_values)

    async def test_export_link_uploads_private_file_and_returns_signed_url(self):
        body = royalty.RoyaltyExportLinkIn(format="xlsx", artist="Artis A")
        result = await royalty.label_royalty_export_link(body, user=self.user)
        self.assertEqual(result["lines"], 1)
        self.assertEqual(result["filename"], "royalti_ArtisA_all.xlsx")
        self.assertTrue(result["url"].startswith("https://r2.example.invalid/report-exports/"))
        (key, (data, content_type)), = self.objects.items()
        self.assertTrue(key.startswith("report-exports/") and key.endswith("/royalti_ArtisA_all.xlsx"))
        self.assertEqual(content_type, report_export.CONTENT_TYPES["xlsx"])
        record = await self.db.report_exports.find_one({})
        self.assertEqual(record["key"], key)
        self.assertGreater(record["expires_epoch"], record["created_epoch"])

    async def test_export_link_explains_when_only_legacy_data_exists(self):
        await self.db.royalty_lines.delete_many({"legacy_settled": {"$ne": True}})
        with self.assertRaises(HTTPException) as raised:
            await royalty.label_royalty_export_link(royalty.RoyaltyExportLinkIn(format="csv"), user=self.user)
        self.assertEqual(raised.exception.status_code, 404)
        self.assertIn("setelah periode legacy", raised.exception.detail)
        self.assertEqual(self.objects, {})

    async def test_expired_report_files_are_removed_and_failed_deletes_retry(self):
        deleted = []

        def delete(*, key):
            if "fail" in key:
                raise RuntimeError("r2 unavailable")
            deleted.append(key)

        await self.db.report_exports.insert_many([
            {"_id": "a" * 32, "key": f"report-exports/{'a' * 32}/x.csv", "expires_epoch": 1},
            {"_id": "b" * 32, "key": f"report-exports/{'b' * 32}/y.csv", "expires_epoch": 10 ** 12},
            {"_id": "c" * 32, "key": "cover/not-an-export.png", "expires_epoch": 1},
            {"_id": "d" * 32, "key": f"report-exports/{'d' * 32}/fail.csv", "expires_epoch": 1},
        ])
        with patch.object(storage_service, "_delete_object_sync", side_effect=delete):
            self.assertEqual(await report_export.cleanup_expired_report_exports(), 1)
        self.assertEqual(deleted, [f"report-exports/{'a' * 32}/x.csv"])
        remaining = {doc["_id"]: doc async for doc in self.db.report_exports.find({})}
        self.assertNotIn("a" * 32, remaining)
        self.assertEqual(remaining["d" * 32]["delete_attempts"], 1)
        self.assertIn("b" * 32, remaining)
        self.assertIn("c" * 32, remaining)

    async def test_workbook_closes_each_artist_sheet_before_the_next(self):
        from openpyxl import Workbook, load_workbook
        await self.db.royalty_lines.insert_many([
            {"label_id": "l1", "period": "2026-04", "status": "available", "artist_name_raw": f"Artis {i:03}",
             "track_title_raw": f"Lagu {i}", "quantity": i, "label_idr": i}
            for i in range(60)
        ])
        open_when_created = []
        original = Workbook.create_sheet

        def tracking_create_sheet(workbook, *args, **kwargs):
            open_when_created.append(sum(1 for ws in workbook.worksheets if not ws.closed))
            return original(workbook, *args, **kwargs)

        base = await royalty._royalty_report_scope(self.user)
        with patch.object(Workbook, "create_sheet", tracking_create_sheet):
            data = await report_export.build_workbook_bytes(base=base, period=None, artist=None)
        self.assertEqual(max(open_when_created), 0)
        wb = load_workbook(io.BytesIO(data))
        self.assertEqual(len(wb.sheetnames), 1 + 2 + 60)
        self.assertEqual(wb["Artis 059"]["C2"].value, "Lagu 59")


class ChatReadTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.db = AsyncMongoMockClient().test
        self.patches = [patch.object(chat, "db", self.db), patch.object(chat, "ensure_chat_indexes")]
        for item in self.patches:
            item.start()
        self.addCleanup(lambda: [item.stop() for item in self.patches])
        chat._cache.clear()
        self.label_user = {"id": "lu", "role": "label", "name": "Label"}
        await self.db.labels.insert_one({"id": "l1", "user_id": "lu", "label_name": "Label"})
        await self.db.chat_conversations.insert_one({"id": "c1", "kind": "support", "label_id": "l1",
                                                     "label_user_id": "lu", "status": "active"})
        start = datetime(2026, 10, 1, tzinfo=timezone.utc)
        await self.db.chat_messages.insert_many([{
            "id": f"m{i}", "conversation_id": "c1", "sender_id": "support" if i % 2 else "lu",
            "body": f"pesan {i}", "created_at": (start + timedelta(minutes=i)).isoformat(),
            "read_by": ["support"] if i % 2 else ["lu"],
        } for i in range(chat.MESSAGE_PAGE + 20)])

    async def test_long_thread_returns_newest_messages_in_order(self):
        with patch.object(chat, "_label_support_status", return_value={"support_online": True, "within_hours": True}):
            result = await chat.label_thread(user=self.label_user)
        ids = [m["id"] for m in result["messages"]]
        self.assertEqual(len(ids), chat.MESSAGE_PAGE)
        self.assertEqual(ids[-1], f"m{chat.MESSAGE_PAGE + 19}")
        self.assertEqual(ids[0], "m20")
        self.assertFalse(result["incremental"])
        unread = await self.db.chat_messages.count_documents({"sender_id": "support", "read_by": {"$ne": "lu"}})
        self.assertEqual(unread, 0)

    async def test_incremental_poll_returns_only_recent_tail(self):
        last = await self.db.chat_messages.find_one({"id": f"m{chat.MESSAGE_PAGE + 19}"})
        with patch.object(chat, "_label_support_status", return_value={"support_online": False, "within_hours": False}):
            result = await chat.label_thread(user=self.label_user, since=last["created_at"])
        self.assertTrue(result["incremental"])
        # One-minute overlap: the newest message plus the one stamped a minute before.
        self.assertEqual([m["id"] for m in result["messages"]], [f"m{chat.MESSAGE_PAGE + 18}", f"m{chat.MESSAGE_PAGE + 19}"])

    async def test_admin_settings_screen_reads_saved_values_not_the_poll_cache(self):
        admin = {"id": "sa", "role": "super_admin", "is_admin": True}
        chat._cache["chat_settings"] = (float("inf"), {**chat.DEFAULT_CHAT_SETTINGS, "auto_reply_message": "lama"})
        await self.db.app_settings.insert_one({"id": "chat_settings", "auto_reply_message": "baru"})
        with patch.object(chat, "_require_admin"), patch.object(chat, "_is_support", return_value=True):
            settings = await chat.get_chat_settings(user=admin)
            self.assertEqual(settings["auto_reply_message"], "baru")
            await chat.update_chat_settings(chat.ChatSettingsIn(auto_reply_message="terbaru"), user=admin)
        self.assertEqual((await chat._get_chat_settings())["auto_reply_message"], "terbaru")

    async def test_unread_counts_are_grouped_per_conversation(self):
        await self.db.chat_messages.insert_one({"id": "x", "conversation_id": "c2", "sender_id": "lu",
                                                "created_at": "2026-10-02T00:00:00+00:00", "read_by": ["lu"]})
        counts = await chat._unread_by_conversation(["c1", "c2"], "support")
        self.assertEqual(counts, {"c1": (chat.MESSAGE_PAGE + 20) // 2, "c2": 1})


class ScheduleSeedThrottleTests(unittest.IsolatedAsyncioTestCase):
    async def test_requests_do_not_write_the_schedule_document_every_time(self):
        import os
        from unittest.mock import AsyncMock
        import serverless_schedule as schedule
        import routes.deps as deps
        db = AsyncMongoMockClient().test
        env = {"VERCEL_ENV": "production", "VERCEL_DEPLOYMENT_ID": "dep-1"}
        with patch.dict(os.environ, env), patch.object(deps, "db_bg", db), \
                patch.object(schedule, "send", new_callable=AsyncMock) as send, \
                patch.object(schedule, "_next_seed_check", 0.0):
            await schedule.seed_schedule("www.rilismusik.com")
            self.assertEqual(send.await_count, 1)
            first = await db.serverless_schedule.find_one({"_id": "active"})
            self.assertEqual(first["deployment"], "dep-1")
            # Within the check interval nothing touches the shared document.
            with patch.object(db.serverless_schedule, "update_one", new_callable=AsyncMock) as write:
                await schedule.seed_schedule("www.rilismusik.com")
                write.assert_not_awaited()
            # After the interval a live tick for this deployment needs no write.
            schedule._next_seed_check = 0.0
            with patch.object(db.serverless_schedule, "update_one", new_callable=AsyncMock) as write:
                await schedule.seed_schedule("www.rilismusik.com")
                write.assert_not_awaited()
            self.assertEqual(send.await_count, 1)
            # An expired tick is claimed again exactly as before.
            await db.serverless_schedule.update_one({"_id": "active"}, {"$set": {"tick_until": 0}})
            schedule._next_seed_check = 0.0
            await schedule.seed_schedule("www.rilismusik.com")
            self.assertEqual(send.await_count, 2)


if __name__ == "__main__":
    unittest.main()
