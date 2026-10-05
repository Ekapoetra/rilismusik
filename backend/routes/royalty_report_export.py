"""Label/artist royalty report files (CSV and Excel).

Reports cover published royalty after the legacy cut-off, including lines that
were already withdrawn. Files are read with the uncapped background client and
handed to the browser through a short-lived private R2 link, so large reports
never pass through the Vercel function response body limit (4.5 MB).
"""
import asyncio
import csv
import io
import re
import time
import uuid
from typing import Any, Dict, Optional

import storage_service
from .deps import db_bg, logger

REPORT_STATUSES = ["pending", "available", "withdrawn"]
STATUS_LABEL = {"pending": "Tertunda", "available": "Tersedia", "withdrawn": "Sudah Dicairkan"}
CONTENT_TYPES = {
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
EXPORT_PREFIX = "report-exports/"
EXPORT_LINK_TTL_SECONDS = 900
EXPORT_RETENTION_SECONDS = 3600
LINE_PROJECTION = {
    "_id": 0, "period": 1, "release_title_raw": 1, "track_title_raw": 1, "artist_name_raw": 1,
    "platform": 1, "country": 1, "isrc": 1, "upc": 1, "quantity": 1, "label_idr": 1, "status": 1,
}
CSV_HEADER = [
    "period", "release_title", "track_title", "artist_name", "platform", "country",
    "isrc", "upc", "streams", "royalty_idr", "status", "sudah_dicairkan",
]


def report_base() -> Dict[str, Any]:
    """Published lines after the legacy cut-off; withdrawn lines stay downloadable."""
    return {"status": {"$in": list(REPORT_STATUSES)}, "legacy_settled": {"$ne": True}}


def report_filename(fmt: str, period: Optional[str], artist: Optional[str]) -> str:
    if fmt == "csv":
        return f"royalty_{period or 'all'}.csv"
    safe = "".join(c for c in (artist or "semua") if c.isalnum() or c in ("-", "_"))[:30] or "semua"
    return f"royalti_{safe}_{period or 'all'}.xlsx"


def safe_sheet_title(name: str, used: set) -> str:
    title = re.sub(r"[\\/*?:\[\]]", " ", str(name or "Artis")).strip()[:28] or "Artis"
    candidate = title
    i = 1
    while candidate.lower() in used:
        suffix = f" {i}"
        candidate = title[: 28 - len(suffix)] + suffix
        i += 1
    used.add(candidate.lower())
    return candidate


def _lines(base: Dict[str, Any], sort: list):
    return db_bg.royalty_lines.find(base, LINE_PROJECTION).sort(sort).allow_disk_use(True)


async def build_csv_bytes(base: Dict[str, Any]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(CSV_HEADER)
    async for it in _lines(base, [("period", -1)]):
        st = it.get("status")
        writer.writerow([
            it.get("period"), it.get("release_title_raw"), it.get("track_title_raw"),
            it.get("artist_name_raw"), it.get("platform"), it.get("country"),
            it.get("isrc"), it.get("upc"), it.get("quantity"), it.get("label_idr"),
            STATUS_LABEL.get(st, st), "Ya" if st == "withdrawn" else "Belum",
        ])
    return buf.getvalue().encode("utf-8")


async def build_workbook_bytes(*, base: Dict[str, Any], period: Optional[str], artist: Optional[str]) -> bytes:
    """Styled workbook: a summary sheet plus one detail sheet per artist.

    Written in openpyxl write-only mode so memory stays flat for labels with a
    long history of (withdrawn) royalty lines. Each write-only sheet streams to
    its own temp file, so a sheet is closed as soon as it is complete; keeping
    one open per artist would exhaust file descriptors on large catalogues.
    """
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.utils import get_column_letter

    PINK, INK, GREEN, ZEBRA = "FF1F8E", "111827", "E7F7EF", "FBF3F8"
    thin = Side(style="thin", color="E5E7EB")
    BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
    HEAD_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    HEAD_FILL = PatternFill("solid", fgColor=PINK)
    ZEBRA_FILL = PatternFill("solid", fgColor=ZEBRA)
    PAID_FILL = PatternFill("solid", fgColor=GREEN)
    TITLE_FONT = Font(name="Calibri", bold=True, size=16, color=INK)
    SUB_FONT = Font(name="Calibri", italic=True, size=10, color="6B7280")
    KEY_FONT = Font(name="Calibri", bold=True, color=INK)
    CENTER = Alignment(horizontal="center", vertical="center")
    IDR_FMT, NUM_FMT = '"Rp"#,##0', '#,##0'
    columns = ["Periode", "Rilisan", "Track", "Artis", "Platform", "Negara",
               "ISRC", "UPC", "Streams", "Royalti IDR", "Status", "Sudah Dicairkan"]
    widths = [11, 26, 26, 22, 14, 12, 16, 16, 12, 15, 16, 15]

    async def aggregate(pipeline):
        return [row async for row in db_bg.royalty_lines.aggregate(pipeline, allowDiskUse=True)]

    overall_rows, artist_rows = await asyncio.gather(
        aggregate([{"$match": base}, {"$group": {"_id": None, "total_idr": {"$sum": "$label_idr"},
                   "total_streams": {"$sum": "$quantity"}, "total_lines": {"$sum": 1}}}]),
        aggregate([{"$match": base}, {"$group": {"_id": "$artist_name_raw", "total_idr": {"$sum": "$label_idr"},
                   "streams": {"$sum": "$quantity"}, "lines": {"$sum": 1}}}, {"$sort": {"total_idr": -1}}]),
    )
    overall = (overall_rows or [{}])[0]

    def styled(sheet, value, *, font=None, fill=None, border=None, fmt=None, align=None):
        c = WriteOnlyCell(sheet, value=value)
        if font: c.font = font
        if fill: c.fill = fill
        if border: c.border = border
        if fmt: c.number_format = fmt
        if align: c.alignment = align
        return c

    def header_row(sheet, labels):
        return [styled(sheet, h, font=HEAD_FONT, fill=HEAD_FILL, border=BORDER, align=CENTER) for h in labels]

    wb = Workbook(write_only=True)
    used_titles = {"ringkasan"}

    def new_detail_sheet(name):
        sheet = wb.create_sheet(title=safe_sheet_title(name, used_titles))
        for j, w in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(j)].width = w
        sheet.freeze_panes = "A2"
        sheet.append(header_row(sheet, columns))
        return sheet

    try:
        # ---------------- Ringkasan ----------------
        ws = wb.create_sheet("Ringkasan")
        for j, w in enumerate([26, 14, 16, 10], start=1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.append([styled(ws, "Laporan Royalti", font=TITLE_FONT)])
        ws.append([styled(ws, "Royalti legacy (sebelum bergabung) tidak termasuk", font=SUB_FONT)])
        ws.append([])
        ws.append([styled(ws, "Periode", font=KEY_FONT), period or "Semua periode"])
        ws.append([styled(ws, "Cakupan", font=KEY_FONT), artist or "Semua Artis"])
        ws.append([])
        for label, key, fmt in (("Total Royalti IDR", "total_idr", IDR_FMT),
                                ("Total Streams", "total_streams", NUM_FMT),
                                ("Total Baris", "total_lines", NUM_FMT)):
            ws.append([styled(ws, label, font=KEY_FONT), styled(ws, overall.get(key) or 0, fmt=fmt)])
        ws.append([])
        ws.append([styled(ws, "Ringkasan per Artis", font=TITLE_FONT)])
        ws.append(header_row(ws, ["Artis", "Streams", "Royalti IDR", "Baris"]))
        for i, row in enumerate(artist_rows):
            fill = ZEBRA_FILL if i % 2 else None
            ws.append([
                styled(ws, row["_id"] or "Tanpa Nama", border=BORDER, fill=fill),
                styled(ws, row.get("streams") or 0, border=BORDER, fill=fill, fmt=NUM_FMT),
                styled(ws, row.get("total_idr") or 0, border=BORDER, fill=fill, fmt=IDR_FMT),
                styled(ws, row.get("lines") or 0, border=BORDER, fill=fill, fmt=NUM_FMT),
            ])
        ws.close()

        # ---------------- Detail per artis ----------------
        # Rows arrive grouped by artist, so each sheet is finished before the next starts.
        sheet, current, rownum = None, object(), 2
        async for it in _lines(base, [("artist_name_raw", 1), ("period", -1)]):
            name = it.get("artist_name_raw") or "Tanpa Nama"
            if name != current:
                if sheet is not None:
                    sheet.close()
                sheet, current, rownum = new_detail_sheet(name), name, 2
            status = it.get("status")
            withdrawn = status == "withdrawn"
            zebra = ZEBRA_FILL if rownum % 2 else None
            values = [it.get("period"), it.get("release_title_raw"), it.get("track_title_raw"),
                      it.get("artist_name_raw"), it.get("platform"), it.get("country"),
                      it.get("isrc"), it.get("upc"), it.get("quantity") or 0,
                      it.get("label_idr") or 0, STATUS_LABEL.get(status, status),
                      "Ya" if withdrawn else "Belum"]
            cells = []
            for j, val in enumerate(values, start=1):
                fill = PAID_FILL if withdrawn and j == 12 else zebra
                cells.append(styled(sheet, val, border=BORDER, fill=fill,
                                    fmt=NUM_FMT if j == 9 else IDR_FMT if j == 10 else None,
                                    align=CENTER if j == 12 else None))
            sheet.append(cells)
            rownum += 1
        if sheet is None and artist:
            new_detail_sheet(artist)

        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()
    except BaseException:
        _discard_workbook(wb)
        raise


def _discard_workbook(wb) -> None:
    """Release the temp files of a write-only workbook that failed to build."""
    for sheet in wb.worksheets:
        try:
            if not sheet.closed:
                sheet.close()
        except Exception:  # noqa: BLE001 - best effort while the original error propagates
            pass
        writer = getattr(sheet, "_writer", None)
        if writer is not None:
            try:
                writer.cleanup()
            except Exception:  # noqa: BLE001
                pass


async def build_report(*, fmt: str, base: Dict[str, Any], period: Optional[str], artist: Optional[str]) -> bytes:
    if fmt == "csv":
        return await build_csv_bytes(base)
    return await build_workbook_bytes(base=base, period=period, artist=artist)


async def publish_report_file(*, actor_id: str, fmt: str, data: bytes, filename: str) -> str:
    """Store a generated report privately and return a short-lived download link."""
    export_id = uuid.uuid4().hex
    key = f"{EXPORT_PREFIX}{export_id}/{filename}"
    await storage_service.upload_bytes(key=key, data=data, content_type=CONTENT_TYPES[fmt])
    now = time.time()
    await db_bg.report_exports.insert_one({
        "_id": export_id, "key": key, "actor_id": actor_id, "format": fmt, "size": len(data),
        "created_epoch": now, "expires_epoch": now + EXPORT_RETENTION_SECONDS,
    })
    return await storage_service.generate_presigned_url(key=key, ttl=EXPORT_LINK_TTL_SECONDS, filename=filename)


async def cleanup_expired_report_exports():
    """Delete generated report files after their retention window.

    A tracking record is removed only after its object is deleted (S3 deletes
    are idempotent, so a missing object also counts), so a failed delete is
    retried on the next run and the collection only holds recent exports.
    """
    removed = 0
    async for document in db_bg.report_exports.find({"expires_epoch": {"$lt": time.time()}}).limit(200):
        key = document.get("key", "")
        if not key.startswith(f"{EXPORT_PREFIX}{document['_id']}/"):
            continue
        try:
            await asyncio.to_thread(storage_service._delete_object_sync, key=key)
        except Exception as exc:  # noqa: BLE001 - keep the record so the next run retries
            await db_bg.report_exports.update_one({"_id": document["_id"]}, {
                "$inc": {"delete_attempts": 1}, "$set": {"last_error": type(exc).__name__},
            })
            continue
        await db_bg.report_exports.delete_one({"_id": document["_id"]})
        removed += 1
    if removed:
        logger.info("[EXPORT] removed %d expired report files", removed)
    return removed
