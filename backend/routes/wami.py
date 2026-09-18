"""WAMI registration orders router."""
from fastapi import APIRouter, HTTPException, Request, Response, Depends, UploadFile, File, Form, Query
from typing import Optional, List, Dict, Any
from datetime import datetime, timezone, timedelta, date
import os
import csv
import io
import shutil
import secrets

from .deps import (
    db, logger, UPLOAD_DIR,
    get_current_user, require_label, require_artist, require_admin, require_super_admin,
    public_user, get_label_by_user, redact_label_for_self, LABEL_HIDDEN_FIELDS,
    log_activity, notify, notify_many, admin_user_ids, label_user_ids,
    LABEL_ROLE, ARTIST_ROLE, ADMIN_ROLES, SUPER_ADMIN,
assert_admin_permission,
)
from models import (
    RegisterLabelIn, LoginIn, ForgotPasswordIn, ResetPasswordIn, VerifyEmailIn,
    LabelProfileUpdate, BankAccountIn,
    ReleaseDraftIn, ReleaseSubmitConfirmation, AdminReleaseAction,
    ArtistIn, ArtistUpdateIn,
    CreateReleasePaymentIn,
    CMSUpdateIn, AdminUserCreateIn, LabelStatusUpdate,
    ExchangeRateIn, RoyaltyImportPublishIn, RoyaltyLineMatchIn,
    WithdrawRequestIn, WithdrawAdminAction,
    TicketCreateIn, TicketCommentIn, TicketAdminUpdateIn,
    ContractCreateIn, ContractExtendIn, ContractTerminateIn,
    BlacklistIn, NotificationMarkIn,
    CreateSubscriptionPaymentIn, CreateWamiOrderIn, AdminWamiUpdateIn,
    WamiMigrationApplyIn,
    now_iso, new_id,
)
from auth_utils import (
    hash_password, verify_password,
    create_access_token, create_refresh_token,
    set_auth_cookies, clear_auth_cookies, decode_token,
)
from royalty_utils import (
    parse_csv_bytes, detect_columns, parse_amount, normalize_header,
    parse_period_from_value, calculate_line, label_percentage_at,
    strip_sensitive,
)
from withdraw_utils import withdraw_window_state, jakarta_now, MIN_WITHDRAW_IDR

# =============================================================================
#                              WAMI ORDERS
# =============================================================================
wami_r = APIRouter(prefix="/wami", tags=["wami"])


@wami_r.get("/label")
async def label_list_wami(user: dict = Depends(require_label)):
    label = await get_label_by_user(user)
    items = await db.wami_orders.find({"label_id": label["id"]}, {"_id": 0}).sort("created_at", -1).to_list(500)
    return items


@wami_r.get("/admin")
async def admin_list_wami(user: dict = Depends(require_admin), status: Optional[str] = None):
    filt: Dict[str, Any] = {}
    if status:
        filt["status"] = status
    items = await db.wami_orders.find(filt, {"_id": 0}).sort("created_at", -1).to_list(500)
    label_ids = list({i["label_id"] for i in items})
    labels = await db.labels.find({"id": {"$in": label_ids}}, {"_id": 0, "id": 1, "label_name": 1, "user_id": 1}).to_list(1000)
    name_map = {lab["id"]: lab["label_name"] for lab in labels}
    for it in items:
        it["label_name"] = name_map.get(it["label_id"])
    return items


@wami_r.post("/admin/{order_id}/status")
async def admin_update_wami(order_id: str, body: AdminWamiUpdateIn, user: dict = Depends(require_admin)):
    assert_admin_permission(user, "wami.manage")
    order = await db.wami_orders.find_one({"id": order_id})
    if not order:
        raise HTTPException(status_code=404, detail="Order tidak ditemukan")
    upd: Dict[str, Any] = {"status": body.status, "updated_at": now_iso()}
    if body.note is not None:
        upd["admin_note"] = body.note
    if body.wami_reference is not None:
        upd["wami_reference"] = body.wami_reference
    if body.status == "registered":
        upd["registered_at"] = now_iso()
    await db.wami_orders.update_one({"id": order_id}, {"$set": upd})
    # Notify label
    user_ids = await label_user_ids(order["label_id"])
    titles = {
        "in_progress": ("WAMI sedang diproses", f"Pendaftaran WAMI '{order.get('track_title')}' sedang diproses."),
        "registered": ("WAMI berhasil terdaftar ✓", f"'{order.get('track_title')}' telah terdaftar di LMKN/WAMI."),
        "rejected": ("WAMI ditolak", f"Pendaftaran '{order.get('track_title')}' ditolak. {body.note or ''}"),
        "cancelled": ("WAMI dibatalkan", f"Pendaftaran '{order.get('track_title')}' dibatalkan."),
    }
    if body.status in titles:
        title, msg = titles[body.status]
        await notify_many(user_ids, f"wami_{body.status}", title, msg, "/label/wami", {"wami_order_id": order_id})
    await log_activity(user["id"], f"wami_{body.status}", "wami", order_id, after={"status": body.status})
    return await db.wami_orders.find_one({"id": order_id}, {"_id": 0})


# =============================================================================
#                       WAMI MIGRATION IMPORT (bulk match)
# =============================================================================
import io as _io
import re as _re
import difflib as _difflib
from collections import defaultdict as _defaultdict
try:
    import openpyxl as _openpyxl
except Exception:  # pragma: no cover
    _openpyxl = None

_WAMI_COLS = ["title", "iswc", "contributors", "original publishers", "performers", "bmat id", "internal id"]


def _wami_norm(value) -> str:
    return _re.sub(r"\s+", " ", str(value or "").strip()).upper()


def _wami_clean(value) -> str:
    v = str(value).strip() if value is not None else ""
    return "" if v == "-" else v


def _wami_split_multi(value):
    raw = _wami_clean(value)
    if not raw:
        return []
    out = []
    for part in _re.split(r"[\n;]+", raw):
        p = part.strip()
        if p and p != "-" and p not in out:
            out.append(p)
    return out


@wami_r.post("/migration/preview")
async def wami_migration_preview(file: UploadFile = File(...), user: dict = Depends(require_admin)):
    """Parse a WAMI migration .xlsx and match each Title to a track/release. Returns a preview;
    nothing is written until /migration/apply is called with the resolved selections."""
    assert_admin_permission(user, "wami.manage")
    if _openpyxl is None:
        raise HTTPException(status_code=500, detail="openpyxl tidak tersedia di server")
    content = await file.read()
    try:
        wb = _openpyxl.load_workbook(_io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise HTTPException(status_code=400, detail="File Excel tidak valid (.xlsx)")
    ws = wb["Works"] if "Works" in wb.sheetnames else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(status_code=400, detail="File kosong")
    header = [_wami_clean(h).lower() for h in rows[0]]
    idx = {c: (header.index(c) if c in header else None) for c in _WAMI_COLS}
    if idx["title"] is None:
        raise HTTPException(status_code=400, detail="Kolom 'Title' tidak ditemukan")

    tracks = await db.tracks.find({}, {"_id": 0, "id": 1, "release_id": 1, "track_title": 1, "performers": 1, "artist_name": 1, "wami": 1}).to_list(200000)
    releases = await db.releases.find({}, {"_id": 0, "id": 1, "release_title": 1, "label_id": 1, "primary_artist": 1}).to_list(200000)
    rel_map = {r["id"]: r for r in releases}
    label_ids = list({r.get("label_id") for r in releases if r.get("label_id")})
    labels = await db.labels.find({"id": {"$in": label_ids}}, {"_id": 0, "id": 1, "label_name": 1}).to_list(100000)
    lab_map = {l["id"]: l.get("label_name") for l in labels}
    tracks_by_release = _defaultdict(list)
    track_by_title = _defaultdict(list)
    for t in tracks:
        tracks_by_release[t.get("release_id")].append(t)
        track_by_title[_wami_norm(t.get("track_title"))].append(t)
    rel_by_title = _defaultdict(list)
    for r in releases:
        rel_by_title[_wami_norm(r.get("release_title"))].append(r)
    norm_track_titles = list(track_by_title.keys())

    def cand(t):
        rel = rel_map.get(t.get("release_id"), {})
        return {"track_id": t["id"], "release_id": t.get("release_id"), "release_title": rel.get("release_title"),
                "track_title": t.get("track_title"), "primary_artist": t.get("artist_name") or rel.get("primary_artist"),
                "label_name": lab_map.get(rel.get("label_id")), "track_performers": t.get("performers"),
                "already_wami": bool(t.get("wami"))}

    out_rows = []
    summary = {"total": 0, "matched": 0, "multi": 0, "not_found": 0, "fuzzy": 0}
    for i, row in enumerate(rows[1:], start=2):
        title = _wami_clean(row[idx["title"]]) if row and idx["title"] < len(row) else ""
        if not title:
            continue
        summary["total"] += 1
        nt = _wami_norm(title)
        match_type = "exact"
        cands = []
        if nt in track_by_title:
            cands = [cand(t) for t in track_by_title[nt]]
        elif nt in rel_by_title:
            for r in rel_by_title[nt]:
                cands += [cand(t) for t in tracks_by_release.get(r["id"], [])]
        else:
            close = _difflib.get_close_matches(nt, norm_track_titles, n=5, cutoff=0.9)
            if close:
                match_type = "fuzzy"
                for c in close:
                    cands += [cand(t) for t in track_by_title[c]]

        def g(col):
            j = idx[col]
            return row[j] if (j is not None and j < len(row)) else None
        wami = {"iswc": _wami_clean(g("iswc")), "contributors": _wami_split_multi(g("contributors")),
                "original_publishers": _wami_split_multi(g("original publishers")), "performers": _wami_split_multi(g("performers")),
                "bmat_id": _wami_clean(g("bmat id")), "internal_id": _wami_clean(g("internal id"))}
        status = "matched" if len(cands) == 1 else ("multi" if len(cands) > 1 else "not_found")
        summary[status if status != "matched" else "matched"] += 1
        if match_type == "fuzzy" and cands:
            summary["fuzzy"] += 1
        out_rows.append({"row": i, "title": title, "wami": wami, "match_type": match_type if cands else None,
                         "status": status, "candidates": cands,
                         "selected_track_id": cands[0]["track_id"] if len(cands) == 1 else None})
    await log_activity(user["id"], "wami_migration_preview", "wami", None, after={"filename": file.filename, **summary})
    return {"summary": summary, "rows": out_rows, "filename": file.filename}


@wami_r.post("/migration/apply")
async def wami_migration_apply(body: WamiMigrationApplyIn, user: dict = Depends(require_admin)):
    """Attach WAMI metadata to the resolved tracks and flag their releases as WAMI-registered."""
    assert_admin_permission(user, "wami.manage")
    applied = 0
    touched = set()
    for it in body.items:
        if not it.track_id:
            continue
        tr = await db.tracks.find_one({"id": it.track_id}, {"_id": 0, "id": 1, "release_id": 1})
        if not tr:
            continue
        wami_block = {"iswc": it.iswc or "", "contributors": it.contributors, "original_publishers": it.original_publishers,
                      "performers": it.performers, "bmat_id": it.bmat_id or "", "internal_id": it.internal_id or "",
                      "source": "migration", "registered_at": now_iso(), "imported_by": user["id"]}
        await db.tracks.update_one({"id": it.track_id}, {"$set": {"wami": wami_block, "wami_registered": True, "updated_at": now_iso()}})
        applied += 1
        rid = tr.get("release_id") or it.release_id
        if rid:
            touched.add(rid)
    for rid in touched:
        cnt = await db.tracks.count_documents({"release_id": rid, "wami_registered": True})
        await db.releases.update_one({"id": rid}, {"$set": {"wami_registered": cnt > 0, "wami_track_count": cnt, "wami_registered_at": now_iso(), "updated_at": now_iso()}})
    batch = {"id": new_id(), "filename": body.filename, "applied_tracks": applied, "releases": len(touched), "by": user["id"], "created_at": now_iso()}
    await db.wami_imports.insert_one(dict(batch))
    await log_activity(user["id"], "wami_migration_apply", "wami", batch["id"], after={"applied": applied, "releases": len(touched)})
    return {"applied": applied, "releases": len(touched), "batch_id": batch["id"]}


@wami_r.get("/migration/history")
async def wami_migration_history(user: dict = Depends(require_admin)):
    assert_admin_permission(user, "wami.manage")
    return await db.wami_imports.find({}, {"_id": 0}).sort("created_at", -1).to_list(100)


