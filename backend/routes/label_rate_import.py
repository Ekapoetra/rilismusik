"""Bulk preview and background synchronization for label royalty rates."""
import asyncio
import csv
import io
from collections import defaultdict
from datetime import datetime, timezone, timedelta
from typing import Any, Dict, List

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel, Field

from models import new_id, now_iso
from royalty_utils import normalize_header, normalize_label_match_name
from .deps import db, db_bg, log_activity, logger, require_admin
from .royalty_recalculation import recalculate_label_unwithdrawn, trigger_royalty_caches


rate_import_r = APIRouter(prefix="/admin/labels/rate-import", tags=["label-rate-import"])
FINANCE_ROLES = ("super_admin", "admin_finance")
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000


class RateImportCommitIn(BaseModel):
    batch_id: str = Field(..., min_length=1, max_length=100)
    reason: str = Field(default="Sinkronisasi bulk rate label", max_length=300)


def _require_finance(user: dict) -> None:
    if user.get("role") not in FINANCE_ROLES:
        raise HTTPException(status_code=403, detail="Hanya Admin Finance / Super Admin")


def _csv_rows(content: bytes) -> List[List[Any]]:
    text = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise HTTPException(status_code=400, detail="Encoding CSV tidak didukung")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
        return [list(row) for row in csv.reader(io.StringIO(text), dialect)]
    except csv.Error:
        return [list(row) for row in csv.reader(io.StringIO(text), delimiter=",")]


def _xlsx_rows(content: bytes) -> List[List[Any]]:
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        rows = []
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            rows.append(list(row))
            if index > MAX_ROWS + 1:
                break
        workbook.close()
        return rows
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"XLSX tidak dapat dibaca: {str(exc)[:180]}") from exc


def _parse_file(filename: str, content: bytes) -> List[List[Any]]:
    lowered = filename.lower()
    if lowered.endswith(".csv"):
        return _csv_rows(content)
    if lowered.endswith(".xlsx"):
        return _xlsx_rows(content)
    raise HTTPException(status_code=400, detail="Format file harus .xlsx atau .csv")


def _parse_rate(value: Any) -> float:
    if value is None or str(value).strip() == "":
        raise ValueError("Rate wajib diisi")
    cleaned = str(value).strip().replace("%", "").replace(",", ".")
    rate = float(cleaned)
    if not 0 <= rate <= 100:
        raise ValueError("Rate harus 0–100")
    return round(rate, 6)


def _extract_input_rows(table: List[List[Any]]) -> List[Dict[str, Any]]:
    nonempty = [row for row in table if any(value not in (None, "") for value in row)]
    if not nonempty:
        raise HTTPException(status_code=400, detail="File kosong")
    headers = [normalize_header(str(value or "")) for value in nonempty[0]]
    try:
        label_index = headers.index("nama label")
        rate_index = headers.index("rate")
    except ValueError as exc:
        raise HTTPException(status_code=400, detail="Kolom wajib: 'Nama Label' dan 'Rate'") from exc
    parsed: List[Dict[str, Any]] = []
    for row_number, row in enumerate(nonempty[1:], start=2):
        label_name = str(row[label_index] if label_index < len(row) and row[label_index] is not None else "").strip()
        rate_raw = row[rate_index] if rate_index < len(row) else None
        parsed.append({"row_number": row_number, "input_label_name": label_name, "rate_raw": rate_raw})
        if len(parsed) > MAX_ROWS:
            raise HTTPException(status_code=400, detail=f"Maksimal {MAX_ROWS} baris per file")
    return parsed


async def _build_preview(input_rows: List[Dict[str, Any]]) -> tuple[List[Dict[str, Any]], Dict[str, int]]:
    labels = await db_bg.labels.find(
        {}, {"_id": 0, "id": 1, "label_name": 1, "royalty_percentage_default": 1},
    ).to_list(10000)
    labels_by_name: Dict[str, List[dict]] = defaultdict(list)
    for label in labels:
        labels_by_name[normalize_label_match_name(label.get("label_name"))].append(label)

    grouped_input: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for row in input_rows:
        grouped_input[normalize_label_match_name(row["input_label_name"])].append(row)
    active_withdraw_ids = set(await db_bg.withdraw_requests.distinct("label_id", {
        "status": {"$in": ["requested", "approved"]}, "legacy_import": {"$ne": True},
    }))

    results: List[Dict[str, Any]] = []
    for normalized_name, grouped_rows in grouped_input.items():
        valid_rates = []
        rate_errors: Dict[int, str] = {}
        for row in grouped_rows:
            try:
                valid_rates.append((row["row_number"], _parse_rate(row["rate_raw"])))
            except (TypeError, ValueError) as exc:
                rate_errors[row["row_number"]] = str(exc)
        unique_rates = {rate for _, rate in valid_rates}
        conflicting = len(unique_rates) > 1
        leader_row_number = min((number for number, _ in valid_rates), default=None)
        candidates = labels_by_name.get(normalized_name, []) if normalized_name else []
        for row in grouped_rows:
            base = {
                "row_number": row["row_number"],
                "input_label_name": row["input_label_name"],
                "normalized_name": normalized_name,
                "input_rate": None,
                "label_id": None,
                "matched_label_name": None,
                "current_rate": None,
                "candidate_labels": [{"id": item["id"], "label_name": item["label_name"]} for item in candidates],
            }
            if row["row_number"] in rate_errors:
                results.append({**base, "status": "invalid", "reason": rate_errors[row["row_number"]]})
                continue
            rate = next(value for number, value in valid_rates if number == row["row_number"])
            base["input_rate"] = rate
            if not normalized_name:
                results.append({**base, "status": "invalid", "reason": "Nama Label wajib diisi"})
            elif conflicting:
                results.append({**base, "status": "duplicate_conflict", "reason": "Nama Label berulang dengan Rate berbeda"})
            elif row["row_number"] != leader_row_number:
                results.append({**base, "status": "duplicate_redundant", "reason": "Duplikat identik—cukup baris pertama yang dipakai"})
            elif not candidates:
                results.append({**base, "status": "unmatched", "reason": "Tidak ada label dengan nama ternormalisasi yang sama"})
            elif len(candidates) > 1:
                results.append({**base, "status": "ambiguous", "reason": "Lebih dari satu label memiliki nama ternormalisasi yang sama"})
            else:
                label = candidates[0]
                stored_rate = label.get("royalty_percentage_default")
                current_rate = float(60 if stored_rate is None else stored_rate)
                matched = {
                    **base,
                    "label_id": label["id"],
                    "matched_label_name": label["label_name"],
                    "current_rate": current_rate,
                    "candidate_labels": [],
                }
                if label["id"] in active_withdraw_ids:
                    results.append({**matched, "status": "blocked_withdraw", "reason": "Ada withdraw requested/approved yang harus diselesaikan lebih dulu"})
                elif current_rate == rate:
                    results.append({**matched, "status": "unchanged", "reason": "Rate sudah sama"})
                else:
                    results.append({**matched, "status": "matched", "reason": "Siap diperbarui dan dihitung ulang"})

    results.sort(key=lambda item: item["row_number"])
    summary = {key: 0 for key in (
        "total_rows", "matched", "unchanged", "unmatched", "ambiguous",
        "invalid", "duplicate_conflict", "duplicate_redundant", "blocked_withdraw",
    )}
    summary["total_rows"] = len(results)
    for row in results:
        summary[row["status"]] += 1
    summary["will_update"] = summary["matched"]
    return results, summary


@rate_import_r.post("/preview")
async def preview_label_rate_import(file: UploadFile = File(...), user: dict = Depends(require_admin)):
    _require_finance(user)
    content = await file.read(MAX_FILE_BYTES + 1)
    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(status_code=400, detail="Ukuran file maksimal 10 MB")
    input_rows = _extract_input_rows(_parse_file(file.filename or "", content))
    rows, summary = await _build_preview(input_rows)
    batch_id = new_id()
    created_at = datetime.now(timezone.utc)
    await db_bg.label_rate_imports.insert_one({
        "id": batch_id,
        "filename": file.filename,
        "status": "preview",
        "summary": summary,
        "rows": rows,
        "submitted_by": user["id"],
        "created_at": created_at.isoformat(),
        "expires_at": (created_at + timedelta(hours=24)).isoformat(),
        "updated_at": created_at.isoformat(),
    })
    return {"batch_id": batch_id, "filename": file.filename, "status": "preview", "summary": summary, "rows": rows}


@rate_import_r.post("/commit")
async def commit_label_rate_import(body: RateImportCommitIn, user: dict = Depends(require_admin)):
    _require_finance(user)
    batch = await db_bg.label_rate_imports.find_one({"id": body.batch_id}, {"_id": 0})
    if not batch:
        raise HTTPException(status_code=404, detail="Preview import tidak ditemukan")
    if batch.get("submitted_by") != user["id"] and user.get("role") != "super_admin":
        raise HTTPException(status_code=403, detail="Preview ini dibuat oleh admin lain")
    if datetime.fromisoformat(batch["expires_at"]) < datetime.now(timezone.utc):
        raise HTTPException(status_code=410, detail="Preview kedaluwarsa—upload ulang file")
    if batch.get("status") in ("queued", "processing", "done", "done_with_errors"):
        return {"batch_id": body.batch_id, "job_id": batch.get("job_id"), "status": batch.get("status"), "already_started": True}
    if not batch.get("summary", {}).get("will_update"):
        raise HTTPException(status_code=400, detail="Tidak ada baris valid yang perlu diperbarui")

    job_id = new_id()
    queued_at = now_iso()
    claimed = await db_bg.label_rate_imports.update_one(
        {"id": body.batch_id, "status": "preview"},
        {"$set": {"status": "queued", "job_id": job_id, "reason": body.reason.strip(), "updated_at": queued_at}},
    )
    if claimed.modified_count != 1:
        current = await db_bg.label_rate_imports.find_one({"id": body.batch_id}, {"_id": 0, "status": 1, "job_id": 1})
        return {"batch_id": body.batch_id, "job_id": (current or {}).get("job_id"), "status": (current or {}).get("status"), "already_started": True}
    await db_bg.migrate_jobs.insert_one({
        "id": job_id,
        "kind": "label_rate_sync",
        "status": "queued",
        "batch_id": body.batch_id,
        "submitted_by": user["id"],
        "submitted_at": queued_at,
        "updated_at": queued_at,
        "progress_labels_done": 0,
        "progress_labels_total": batch["summary"]["will_update"],
    })
    asyncio.create_task(_run_rate_sync_job(
        batch_id=body.batch_id, job_id=job_id, user_id=user["id"], reason=body.reason.strip(),
    ))
    return {"batch_id": body.batch_id, "job_id": job_id, "status": "queued", "already_started": False}


@rate_import_r.get("/jobs/{job_id}")
async def get_label_rate_import_job(job_id: str, user: dict = Depends(require_admin)):
    _require_finance(user)
    job = await db.migrate_jobs.find_one({"id": job_id, "kind": "label_rate_sync"}, {"_id": 0})
    if not job:
        raise HTTPException(status_code=404, detail="Job sinkronisasi tidak ditemukan")
    return job


async def _set_row_result(batch_id: str, row_number: int, **values: Any) -> None:
    await db_bg.label_rate_imports.update_one(
        {"id": batch_id},
        {"$set": {f"rows.$[target].{key}": value for key, value in values.items()}},
        array_filters=[{"target.row_number": row_number}],
    )


async def _run_rate_sync_job(*, batch_id: str, job_id: str, user_id: str, reason: str) -> None:
    totals = {
        "labels_updated": 0,
        "labels_recalculated": 0,
        "labels_skipped_active_withdraw": 0,
        "labels_skipped_changed_after_preview": 0,
        "rows_failed": 0,
        "lines_recalculated": 0,
        "pending_delta_idr": 0,
        "available_delta_idr": 0,
    }
    try:
        batch = await db_bg.label_rate_imports.find_one({"id": batch_id}, {"_id": 0})
        rows = [
            row for row in (batch or {}).get("rows", [])
            if row.get("status") == "matched" and row.get("apply_status") != "done"
        ]
        await db_bg.label_rate_imports.update_one({"id": batch_id}, {"$set": {"status": "processing", "updated_at": now_iso()}})
        await db_bg.migrate_jobs.update_one({"id": job_id}, {"$set": {"status": "processing", "phase": "updating_labels", "started_at": now_iso(), "updated_at": now_iso()}})
        for index, row in enumerate(rows, start=1):
            label_id = row["label_id"]
            try:
                active = await db_bg.withdraw_requests.find_one({
                    "label_id": label_id,
                    "status": {"$in": ["requested", "approved"]},
                    "legacy_import": {"$ne": True},
                }, {"_id": 0, "id": 1})
                if active:
                    totals["labels_skipped_active_withdraw"] += 1
                    await _set_row_result(batch_id, row["row_number"], apply_status="blocked_withdraw", apply_reason="Withdraw aktif muncul setelah preview")
                    continue
                label = await db_bg.labels.find_one(
                    {"id": label_id},
                    {
                        "_id": 0,
                        "royalty_percentage_default": 1,
                        "royalty_recalculation_status": 1,
                        "royalty_recalculation_job_id": 1,
                        "label_name": 1,
                    },
                )
                if not label:
                    raise RuntimeError("Label sudah tidak ditemukan")
                stored_rate = label.get("royalty_percentage_default")
                current_rate = float(60 if stored_rate is None else stored_rate)
                new_rate = float(row["input_rate"])
                resume_same_job = (
                    current_rate == new_rate
                    and label.get("royalty_recalculation_status") == "processing"
                    and label.get("royalty_recalculation_job_id") == job_id
                )
                if current_rate != float(row["current_rate"]) and not resume_same_job:
                    totals["labels_skipped_changed_after_preview"] += 1
                    await _set_row_result(batch_id, row["row_number"], apply_status="stale_conflict", apply_reason=f"Rate berubah setelah preview ({current_rate:g}%)")
                    continue
                changed_at = now_iso()
                if not resume_same_job:
                    rate_filter: Dict[str, Any] = {
                        "id": label_id,
                        "royalty_recalculation_status": {"$ne": "processing"},
                    }
                    if stored_rate is None:
                        rate_filter["$or"] = [
                            {"royalty_percentage_default": {"$exists": False}},
                            {"royalty_percentage_default": None},
                        ]
                    else:
                        rate_filter["royalty_percentage_default"] = stored_rate
                    claimed = await db_bg.labels.update_one(
                        rate_filter,
                        {"$set": {
                            "royalty_percentage_default": new_rate,
                            "royalty_recalculation_status": "processing",
                            "royalty_recalculation_job_id": job_id,
                            "updated_at": changed_at,
                        }},
                    )
                    if claimed.matched_count != 1:
                        totals["labels_skipped_changed_after_preview"] += 1
                        await _set_row_result(batch_id, row["row_number"], apply_status="stale_conflict", apply_reason="Label sedang diproses atau Rate berubah setelah preview")
                        continue
                await db_bg.royalty_percentage_history.update_one(
                    {"id": f"label-rate-sync:{batch_id}:{label_id}"},
                    {"$setOnInsert": {
                        "id": f"label-rate-sync:{batch_id}:{label_id}",
                        "label_id": label_id,
                        "percentage": new_rate,
                        "effective_month": datetime.now(timezone.utc).strftime("%Y-%m"),
                        "changed_by": user_id,
                        "changed_at": changed_at,
                        "reason": reason or "Sinkronisasi bulk rate label",
                        "batch_id": batch_id,
                    }},
                    upsert=True,
                )
                if not resume_same_job:
                    totals["labels_updated"] += 1
                recalculated = await recalculate_label_unwithdrawn(
                    label_id=label_id,
                    percentage=new_rate,
                    job_id=job_id,
                )
                totals["labels_recalculated"] += 1
                for key in ("lines_recalculated", "pending_delta_idr", "available_delta_idr"):
                    totals[key] += int(recalculated.get(key) or 0)
                await _set_row_result(batch_id, row["row_number"], apply_status="done", apply_reason="Rate dan royalti unsettled selesai disinkronkan", recalculation=recalculated)
            except Exception as exc:
                totals["rows_failed"] += 1
                logger.exception("[LABEL RATE SYNC] batch=%s label=%s failed: %s", batch_id, label_id, exc)
                await _set_row_result(batch_id, row["row_number"], apply_status="error", apply_reason=f"{type(exc).__name__}: {str(exc)[:240]}")
            finally:
                await db_bg.migrate_jobs.update_one(
                    {"id": job_id},
                    {"$set": {"progress_labels_done": index, "phase": "recalculating", "updated_at": now_iso()}},
                )

        final_status = "done_with_errors" if totals["rows_failed"] else "done"
        finished_at = now_iso()
        await db_bg.label_rate_imports.update_one(
            {"id": batch_id},
            {"$set": {"status": final_status, "result": totals, "finished_at": finished_at, "updated_at": finished_at}},
        )
        await db_bg.migrate_jobs.update_one(
            {"id": job_id},
            {"$set": {"status": final_status, "phase": "done", "result": totals, "finished_at": finished_at, "updated_at": finished_at}},
        )
        await log_activity(user_id, "bulk_sync_label_rates", "label", None, after={"batch_id": batch_id, **totals})
        await trigger_royalty_caches()
    except Exception as exc:
        logger.exception("[LABEL RATE SYNC] job=%s failed: %s", job_id, exc)
        failed_at = now_iso()
        error_message = f"{type(exc).__name__}: {str(exc)[:400]}"
        await db_bg.label_rate_imports.update_one({"id": batch_id}, {"$set": {"status": "error", "error_message": error_message, "updated_at": failed_at}})
        await db_bg.migrate_jobs.update_one({"id": job_id}, {"$set": {"status": "error", "error_message": error_message, "finished_at": failed_at, "updated_at": failed_at}})


async def resume_label_rate_jobs() -> None:
    """Resume rate-sync jobs interrupted by hot reload or pod replacement."""
    jobs = await db_bg.migrate_jobs.find(
        {"kind": "label_rate_sync", "status": {"$in": ["queued", "processing"]}},
        {"_id": 0, "id": 1, "batch_id": 1, "submitted_by": 1},
    ).to_list(20)
    for job in jobs:
        batch = await db_bg.label_rate_imports.find_one(
            {"id": job.get("batch_id")}, {"_id": 0, "reason": 1},
        )
        if not batch:
            await db_bg.migrate_jobs.update_one(
                {"id": job["id"]},
                {"$set": {"status": "error", "error_message": "Batch preview tidak ditemukan saat resume", "updated_at": now_iso()}},
            )
            continue
        asyncio.create_task(_run_rate_sync_job(
            batch_id=job["batch_id"],
            job_id=job["id"],
            user_id=job.get("submitted_by") or "system",
            reason=batch.get("reason") or "Resume sinkronisasi bulk rate label",
        ))